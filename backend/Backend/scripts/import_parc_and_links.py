from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.enums import EquipmentStatus
from app.db.session import SessionLocal
from app.models.ai_core import EquipmentHistoricalLink
from app.models.equipment import Equipment, EquipmentCategory
from app.models.historical import HistoricalRequest


CATEGORY_MAP: dict[str, tuple[str, str]] = {
    "PC": ("PC", "Poste fixe"),
    "PORTABLE": ("PORTABLE", "Ordinateur portable"),
    "ALL-IN-ONE": ("ALL_IN_ONE", "Ordinateur tout-en-un"),
    "LASER MONO": ("IMPRIMANTE", "Imprimante laser monochrome"),
    "LASER COULEUR": ("IMPRIMANTE", "Imprimante laser couleur"),
    "LASER COULEUR A3": ("IMPRIMANTE", "Imprimante laser couleur A3"),
    "MATRICIELLE": ("IMPRIMANTE", "Imprimante matricielle"),
    "JET ENCRE": ("IMPRIMANTE", "Imprimante jet d'encre"),
    "TRACEUR": ("IMPRIMANTE", "Traceur"),
    "SCANNER": ("SCANNER", "Scanner"),
    "SCANNER A3": ("SCANNER", "Scanner A3"),
    "SERVEUR": ("SERVEUR", "Serveur"),
    "VIDEO PROJECTEUR": ("VIDEO_PROJECTEUR", "Vidéoprojecteur"),
    "ECRAN TFT": ("ECRAN", "Écran"),
    "LIGNE": ("LIGNE", "Ligne réseau"),
}


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_text(value: Any) -> str:
    text = clean(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(
        character
        for character in text
        if not unicodedata.combining(character)
    )
    return text.upper()


def normalize_header(value: Any) -> str:
    text = normalize_text(value).lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def category_code_and_name(raw_category: str) -> tuple[str, str]:
    key = normalize_text(raw_category)

    if key in CATEGORY_MAP:
        return CATEGORY_MAP[key]

    safe_code = re.sub(r"[^A-Z0-9]+", "_", key).strip("_")
    return safe_code[:50] or "AUTRE", clean(raw_category) or "Autre"


def read_parc_rows(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )
    sheet = workbook.active

    iterator = sheet.iter_rows(values_only=True)
    headers = [normalize_header(value) for value in next(iterator)]
    rows: list[dict[str, str]] = []

    for values in iterator:
        raw = {
            headers[index]: clean(value)
            for index, value in enumerate(values)
            if index < len(headers)
        }

        code = raw.get("code", "").upper()
        serial = (
            raw.get("ns", "")
            or raw.get("numero_serie", "")
        ).upper()

        if not code:
            continue

        rows.append(
            {
                "matricule": raw.get("matricule", ""),
                "categorie": raw.get("categorie", ""),
                "marque": raw.get("marque", ""),
                "modele": raw.get("modele", ""),
                "code": code,
                "serial_number": serial,
            }
        )

    workbook.close()
    return rows


def read_join_rows(path: Path) -> list[dict[str, str]]:
    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
        errors="replace",
    ) as file:
        reader = csv.DictReader(file, delimiter=";")
        return [
            {
                key: clean(value)
                for key, value in row.items()
            }
            for row in reader
        ]


def get_or_create_category(
    db,
    raw_category: str,
    category_cache: dict[str, EquipmentCategory],
) -> EquipmentCategory:
    code, name = category_code_and_name(raw_category)

    if code in category_cache:
        return category_cache[code]

    category = db.scalar(
        select(EquipmentCategory).where(
            EquipmentCategory.code == code
        )
    )

    if category is None:
        category = EquipmentCategory(
            code=code,
            name=name,
            description="Catégorie importée depuis Parc.xlsx",
        )
        db.add(category)
        db.flush()

    category_cache[code] = category
    return category


def import_parc(
    db,
    rows: list[dict[str, str]],
    overwrite: bool,
) -> Counter:
    stats = Counter()
    category_cache: dict[str, EquipmentCategory] = {}

    existing_equipment = {
        equipment.code.upper(): equipment
        for equipment in db.scalars(
            select(Equipment)
        ).all()
    }
    serial_owner = {
        equipment.serial_number.upper(): equipment
        for equipment in existing_equipment.values()
        if equipment.serial_number
    }

    for index, row in enumerate(rows, start=1):
        category = get_or_create_category(
            db,
            row["categorie"],
            category_cache,
        )

        equipment = existing_equipment.get(row["code"])
        serial = row["serial_number"] or None

        if equipment is None:
            if serial and serial in serial_owner:
                stats["serial_conflicts"] += 1
                continue

            equipment = Equipment(
                code=row["code"],
                category_id=category.id,
                brand=row["marque"] or None,
                model=row["modele"] or None,
                serial_number=serial,
                status=EquipmentStatus.IN_SERVICE.value,
                notes=(
                    "Import Parc.xlsx"
                    + (
                        f" | Matricule source: {row['matricule']}"
                        if row["matricule"]
                        else ""
                    )
                ),
            )
            db.add(equipment)
            db.flush()
            existing_equipment[row["code"]] = equipment
            if serial:
                serial_owner[serial] = equipment
            stats["created"] += 1
        else:
            changed = False

            if equipment.category_id != category.id:
                equipment.category_id = category.id
                changed = True

            for attribute, new_value in (
                ("brand", row["marque"] or None),
                ("model", row["modele"] or None),
            ):
                current_value = getattr(equipment, attribute)
                if overwrite or not current_value:
                    if current_value != new_value:
                        setattr(equipment, attribute, new_value)
                        changed = True

            if serial and (overwrite or not equipment.serial_number):
                owner = serial_owner.get(serial)
                if owner is None or owner.id == equipment.id:
                    if equipment.serial_number != serial:
                        equipment.serial_number = serial
                        serial_owner[serial] = equipment
                        changed = True
                else:
                    stats["serial_conflicts"] += 1

            if changed:
                stats["updated"] += 1
            else:
                stats["unchanged"] += 1

        if index % 500 == 0:
            db.flush()
            print(f"[PARC] {index}/{len(rows)}")

    db.flush()
    return stats


def import_links(
    db,
    rows: list[dict[str, str]],
) -> Counter:
    stats = Counter()

    request_numbers = {
        row.get("numero_demande", "")
        for row in rows
        if row.get("statut_jointure") == "ASSOCIE_AUTO"
        and row.get("numero_demande")
    }
    equipment_codes = {
        row.get("parc_code", "").upper()
        for row in rows
        if row.get("statut_jointure") == "ASSOCIE_AUTO"
        and row.get("parc_code")
    }

    request_map = {
        request.numero_demande: request
        for request in db.scalars(
            select(HistoricalRequest).where(
                HistoricalRequest.numero_demande.in_(request_numbers)
            )
        ).all()
    }
    equipment_map = {
        equipment.code.upper(): equipment
        for equipment in db.scalars(
            select(Equipment).where(
                Equipment.code.in_(equipment_codes)
            )
        ).all()
    }

    existing_pairs = {
        (link.equipment_id, link.historical_request_id): link
        for link in db.scalars(
            select(EquipmentHistoricalLink)
        ).all()
    }

    for index, row in enumerate(rows, start=1):
        if row.get("statut_jointure") != "ASSOCIE_AUTO":
            stats["ignored_non_auto"] += 1
            continue

        request_number = row.get("numero_demande", "")
        equipment_code = row.get("parc_code", "").upper()

        request = request_map.get(request_number)
        equipment = equipment_map.get(equipment_code)

        if request is None:
            stats["missing_request"] += 1
            continue
        if equipment is None:
            stats["missing_equipment"] += 1
            continue

        pair = (equipment.id, request.id)
        link = existing_pairs.get(pair)

        try:
            score = float(row.get("score_confiance", "0") or 0)
        except ValueError:
            score = 0.0

        if link is None:
            link = EquipmentHistoricalLink(
                equipment_id=equipment.id,
                historical_request_id=request.id,
                link_method=(
                    row.get("methode_jointure")
                    or "IMPORT_JOINTURE"
                ),
                confidence_score=score,
                validation_status="AUTO",
                source_reference=request_number,
                notes="Import depuis khadamate_parc_jointure.csv",
            )
            db.add(link)
            existing_pairs[pair] = link
            stats["created"] += 1
        else:
            link.link_method = (
                row.get("methode_jointure")
                or link.link_method
            )
            link.confidence_score = score
            link.validation_status = "AUTO"
            stats["updated"] += 1

        if index % 500 == 0:
            db.flush()
            print(f"[LIENS] {index}/{len(rows)}")

    db.flush()
    return stats


def print_stats(title: str, stats: Counter) -> None:
    print(f"\n{title}")
    print("-" * len(title))
    if not stats:
        print("Aucun résultat")
        return

    for key, value in sorted(stats.items()):
        print(f"{key}: {value}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Importe Parc.xlsx puis les liens fiables "
            "Parc × Khadamate."
        )
    )
    parser.add_argument("--parc", required=True)
    parser.add_argument("--jointure", required=True)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument(
        "--commit",
        action="store_true",
        help=(
            "Valide les changements. Sans cette option, "
            "un rollback est effectué."
        ),
    )
    args = parser.parse_args()

    parc_path = Path(args.parc).expanduser().resolve()
    jointure_path = Path(args.jointure).expanduser().resolve()

    if not parc_path.exists():
        raise FileNotFoundError(
            f"Parc introuvable : {parc_path}"
        )
    if not jointure_path.exists():
        raise FileNotFoundError(
            f"Jointure introuvable : {jointure_path}"
        )

    parc_rows = read_parc_rows(parc_path)
    join_rows = read_join_rows(jointure_path)

    print(f"Parc lu : {len(parc_rows)} équipement(s)")
    print(f"Jointure lue : {len(join_rows)} demande(s)")

    with SessionLocal() as db:
        parc_stats = import_parc(
            db,
            parc_rows,
            overwrite=args.overwrite,
        )
        link_stats = import_links(db, join_rows)

        print_stats("RÉSULTAT PARC", parc_stats)
        print_stats("RÉSULTAT LIENS", link_stats)

        if args.commit:
            db.commit()
            print("\nIMPORT VALIDÉ DANS POSTGRESQL")
        else:
            db.rollback()
            print(
                "\nMODE TEST : aucune modification conservée."
            )


if __name__ == "__main__":
    main()
